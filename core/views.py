from django.shortcuts import render
from django.contrib.auth import authenticate, login,logout
from django.shortcuts import redirect
from .models import *
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
def index(request):
    if request.method == "POST":
        username = request.POST.get("username")
        password = request.POST.get("password")
        user=authenticate(request,username=username,password=password)
        if user is not None:
            login(request,user)
            profile = UserProfile.objects.get(user=user)
            if profile.role == "evaluator":
                return redirect("home")
            else:
                return redirect("upload")
        else:
            return render(request,"index.html",{"error":"Invalid username or password"})
    return render(request,"index.html")
def download_data(request, branch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Research Papers Data"
    ws.merge_cells("A1:I1")
    ws["A1"] = f"Research Papers Data - {branch}"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = [
        "Roll Number","Research Paper","Branch","Abstract",
        "Methodology","Results","Formatting","Conclusion","Overall Score"
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in ws[2]:
        col.font = header_font
        col.fill = header_fill
        col.alignment = Alignment(horizontal="center")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    for paper in Uploads.objects.filter(branch=branch):
        ws.append([
            paper.roll_number,
            paper.file.name,
            paper.branch,
            paper.abstract,
            paper.research_methodology,
            paper.results,
            paper.formatting,
            paper.conclusion,
            paper.overall_score
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
    widths = [20,40,15,10,15,10,10,10,15]
    for i,width in enumerate(widths):
        ws.column_dimensions[chr(65+i)].width = width

    response = HttpResponse(content_type="application/ms-excel")
    response["Content-Disposition"] = f'attachment; filename="{branch}_research_data.xlsx"'
    wb.save(response)

    return response
def home(request):
    profile = UserProfile.objects.get(user=request.user)
    if profile.role != "evaluator":
        return HttpResponse("Only Evaluators can access this page")
    papers = Uploads.objects.filter(branch=profile.branch).exclude(id__in=Evaluator.objects.filter(evaluator=request.user).values_list('papers', flat=True))
    selected_paper = None
    if request.method == "POST":
        paper_id = request.POST.get("paper_id")
        if paper_id:
            selected_paper = Uploads.objects.get(id=paper_id)
        if request.POST.get("abstract"):
            abstract = int(request.POST.get("abstract", 0))
            methodology = int(request.POST.get("methodology", 0))
            results = int(request.POST.get("results", 0))
            formatting = int(request.POST.get("formatting", 0))
            conclusion = int(request.POST.get("conclusion", 0))
            abstract_score = abstract * 1
            methodology_score = methodology * 2
            results_score = results * 2
            formatting_score = formatting * 0.5
            conclusion_score = conclusion * 0.5
            total = (
                abstract_score +
                methodology_score +
                results_score +
                formatting_score +
                conclusion_score
            )
            if Evaluator.objects.filter(evaluator=request.user):
                evaluator=Evaluator.objects.get(evaluator=request.user)
                evaluator.papers.add(selected_paper)
                evaluator.save()
            else:
                evaluator=Evaluator.objects.create(
                evaluator=request.user,
                name=request.POST.get("name"),
                employee_id=request.POST.get("id"),
                designation=request.POST.get("designation")
            )
                evaluator.papers.add(selected_paper)
                evaluator.save()
            count = selected_paper.evaluations.count()
            if count == 1:
                selected_paper.abstract = abstract_score
                selected_paper.research_methodology = methodology_score
                selected_paper.results = results_score
                selected_paper.formatting = formatting_score
                selected_paper.conclusion = conclusion_score
                selected_paper.overall_score = total
            else:
                selected_paper.abstract = (selected_paper.abstract + abstract_score) / 2
                selected_paper.research_methodology = (selected_paper.research_methodology + methodology_score) / 2
                selected_paper.results = (selected_paper.results + results_score) / 2
                selected_paper.formatting = (selected_paper.formatting + formatting_score) / 2
                selected_paper.conclusion = (selected_paper.conclusion + conclusion_score) / 2
                selected_paper.overall_score = (selected_paper.overall_score + total) / 2
            selected_paper.save()
    return render(request, "home.html", {
        "papers": papers,
        "selected_paper": selected_paper,
    })
def logout_app(request):
    logout(request)
    return redirect("index")
def upload(request):
    profile = UserProfile.objects.get(user=request.user)
    if profile.role != "dean":
        return HttpResponse("Only dean can upload papers")
    if request.method == "POST":
        file = request.FILES.get("file")
        roll_number=request.POST.get("roll_number")
        branch = request.POST.get("branch")
        Uploads.objects.create(
            file=file,
            branch=branch,
            roll_number=roll_number
        )
    branches=["CSE","ECE",'MECH','CIVIL','EEE','IT','CSM','CSD']
    papers_data={}
    total_papers = 0
    evaluated1_count = 0
    evaluated2_count = 0
    total_evaluated=0
    for branch in branches:
        papers = Uploads.objects.filter(branch=branch)
        top_papers = papers.order_by('-overall_score')[:5]
        count = papers.count()
        evaluated1_count = Evaluator.objects.filter(evaluator__username=f"{branch}_evaluator1",papers__branch=branch).values("papers").distinct().count()
        evaluated2_count = Evaluator.objects.filter(evaluator__username=f"{branch}_evaluator2",papers__branch=branch).values("papers").distinct().count()
        papers_data[branch] = {
        "total": count,
        "evaluated1": evaluated1_count,
        "evaluated2": evaluated2_count,
        "top_papers": top_papers
    }

        total_papers += count
        total_evaluated += evaluated1_count + evaluated2_count
    return render(request,"upload.html",{"papers_data":papers_data,"total_papers":total_papers,"total_evaluated":total_evaluated})
